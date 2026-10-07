import io
from datetime import timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.crawlers.base import NoteData
from sqlalchemy import select

from app.db import session_scope, utcnow
from app.jobs import pipeline
from app.main import app
from app.config import get_settings
from app.main import seed
from app.models import CrawlState, KeywordGroup



@pytest.fixture()
def client(env):
    with TestClient(app) as c:
        yield c


def test_no_login_needed(client):
    assert client.get("/api/v1/hits").status_code == 200
    assert client.get("/api/v1/auth/me").status_code == 404
    assert client.get("/api/v1/users").status_code == 404


def test_default_keyword_seeded_once(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "default_keywords", "全嘻嘻")
    seed()
    groups = client.get("/api/v1/keyword-groups").json()["items"]
    assert [(g["name"], g["words"], g["platforms"]) for g in groups] == [("全嘻嘻", ["全嘻嘻"], ["xhs"])]
    client.delete(f"/api/v1/keyword-groups/{groups[0]['id']}")
    seed()  # 删掉后重启不会再加回来
    assert client.get("/api/v1/keyword-groups").json()["items"] == []


def test_default_keyword_not_added_when_groups_exist(client, monkeypatch):
    with session_scope() as db:
        db.add(KeywordGroup(name="已有", words=["a"], exclude_words=[], platforms=["xhs"]))
    monkeypatch.setattr(get_settings(), "default_keywords", "全嘻嘻")
    seed()
    with session_scope() as db:
        assert list(db.scalars(select(KeywordGroup.name))) == ["已有"]


def test_keyword_group_crud_and_limits(client):
    r = client.post("/api/v1/keyword-groups", json={"name": "品牌词", "words": ["山岚咖啡", " 山岚咖啡 ", "山岚"],
                                                     "exclude_words": ["山岚雾气"]})
    assert r.status_code == 201 and r.json()["words"] == ["山岚咖啡", "山岚"]
    gid = r.json()["id"]
    assert client.post("/api/v1/keyword-groups", json={"name": "x", "words": ["a"], "platforms": ["dy"]}).json()["message"] == "V1 仅支持小红书"
    assert client.post("/api/v1/keyword-groups", json={"name": "x", "words": [str(i) for i in range(21)]}).status_code == 400
    assert client.post("/api/v1/keyword-groups", json={"name": "x", "words": []}).status_code == 400
    client.post("/api/v1/keyword-groups", json={"name": "b", "words": ["b"]})
    client.post("/api/v1/keyword-groups", json={"name": "c", "words": ["c"]})
    r = client.post("/api/v1/keyword-groups", json={"name": "d", "words": ["d"]})
    assert r.status_code == 400 and "最多 3 个" in r.json()["message"]
    r = client.put(f"/api/v1/keyword-groups/{gid}", json={"enabled": False})
    assert r.json()["enabled"] is False
    assert client.delete(f"/api/v1/keyword-groups/{gid}").status_code == 200
    assert len(client.get("/api/v1/keyword-groups").json()["items"]) == 2


def seed_hits(env):
    with session_scope() as db:
        db.add(CrawlState(platform="xhs", keyword="冷萃挂耳", backfill_done=True))
    now = utcnow()
    env.crawler.pages["冷萃挂耳"] = [[
        NoteData(platform_post_id="a1", title="避雷冷萃挂耳", content="太失望了", author_name="小A", published_at=now - timedelta(minutes=3)),
        NoteData(platform_post_id="a2", title="冷萃挂耳回购", content="推荐", author_name="小B", published_at=now - timedelta(minutes=2)),
    ]]
    pipeline.run_due(env.services)


def test_hits_list_filter_patch_export(client, env):
    client.post("/api/v1/keyword-groups", json={"name": "产品词", "words": ["冷萃挂耳"]})
    seed_hits(env)
    items = client.get("/api/v1/hits").json()["items"]
    assert len(items) == 2
    neg = client.get("/api/v1/hits", params={"sentiment": "neg"}).json()["items"]
    assert len(neg) == 1 and neg[0]["post"]["author_name"] == "小A"
    assert client.get("/api/v1/hits", params={"q": "小B"}).json()["items"][0]["sentiment"] == "pos"
    page1 = client.get("/api/v1/hits", params={"limit": 1}).json()
    page2 = client.get("/api/v1/hits", params={"limit": 1, "cursor": page1["next_cursor"]}).json()
    assert page1["next_cursor"] and len(page2["items"]) == 1 and page2["next_cursor"] is None

    hid = neg[0]["id"]
    r = client.patch(f"/api/v1/hits/{hid}", json={"sentiment_manual": "neu"})
    assert r.json()["sentiment"] == "neu" and r.json()["manual"] is True
    client.patch(f"/api/v1/hits/{hid}", json={"status": "done"})
    assert len(client.get("/api/v1/hits").json()["items"]) == 1
    assert len(client.get("/api/v1/hits", params={"status": "all"}).json()["items"]) == 2
    client.patch(f"/api/v1/hits/{hid}", json={"is_irrelevant": True})
    assert len(client.get("/api/v1/hits", params={"status": "all"}).json()["items"]) == 1

    s = client.get("/api/v1/summary").json()
    assert s["today_hits"] == 1 and s["system"]["crawler_mode"] == "none"

    r = client.get("/api/v1/hits/export")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 2 and ws["A1"].value == "发布时间"


def test_alert_detail_and_feedback(client, env):
    client.put("/api/v1/settings/notify", json={"wecom_enabled": True,
                                                "wecom_webhook": "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=k"})
    client.post("/api/v1/keyword-groups", json={"name": "产品词", "words": ["冷萃挂耳"]})
    seed_hits(env)
    neg = client.get("/api/v1/hits", params={"sentiment": "neg"}).json()["items"][0]
    assert neg["alert_id"]
    a = client.get(f"/api/v1/alerts/{neg['alert_id']}").json()
    assert a["deliveries"][0]["channel"] == "wecom" and a["deliveries"][0]["status"] == "ok"
    assert client.patch(f"/api/v1/alerts/{neg['alert_id']}", json={"feedback": "valid"}).json()["feedback"] == "valid"


def test_notify_settings_validation_and_test_push(client, env):
    r = client.put("/api/v1/settings/notify", json={"wecom_webhook": "https://evil.example.com/hook"})
    assert r.status_code == 400
    r = client.put("/api/v1/settings/notify", json={"feishu_enabled": True})
    assert r.status_code == 400 and "飞书" in r.json()["message"]
    r = client.put("/api/v1/settings/notify", json={"quiet_start": "25:00"})
    assert r.status_code == 422
    r = client.put("/api/v1/settings/notify", json={
        "feishu_enabled": True, "feishu_webhook": "https://open.feishu.cn/open-apis/bot/v2/hook/abc",
        "quiet_enabled": True, "quiet_start": "23:00", "quiet_end": "08:00"})
    assert r.status_code == 200 and r.json()["feishu_enabled"] is True
    assert client.post("/api/v1/settings/notify/test").json()["result"] == {"feishu": "ok"}


def test_healthz(client):
    assert client.get("/healthz").json()["ok"] is True
